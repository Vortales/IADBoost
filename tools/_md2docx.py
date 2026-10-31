# -*- coding: utf-8 -*-
"""Convert the legal markdown drafts in AntiGubrnator/ to formatted DOCX."""
import re, sys
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BODY_FONT = "Times New Roman"
INLINE = re.compile(r"(\*\*.+?\*\*|\*[^*]+?\*|`[^`]+?`)")


def add_runs(par, text, bold=False, italic=False, size=None):
    text = text.replace("\\_", "_").replace("\\*", "*").replace("\\|", "|")
    for part in INLINE.split(text):
        if not part:
            continue
        b, i, mono = bold, italic, False
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            part, b = part[2:-2], True
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            part, i = part[1:-1], True
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            part, mono = part[1:-1], True
        run = par.add_run(part)
        run.bold = b
        run.italic = i
        run.font.name = BODY_FONT
        run._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
        if size:
            run.font.size = Pt(size)
        if mono:
            run.font.name = "Consolas"


def set_cell_border(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "808080")
        borders.append(el)
    tcPr.append(borders)


def shade(cell, color="EDEDED"):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    tcPr.append(shd)


def para(doc, text, size=12, bold=False, italic=False, align="justify",
         space_before=0, space_after=6, indent=1.25, left=0, line=1.15,
         keep_with_next=False):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.alignment = {"justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
                    "center": WD_ALIGN_PARAGRAPH.CENTER,
                    "left": WD_ALIGN_PARAGRAPH.LEFT,
                    "right": WD_ALIGN_PARAGRAPH.RIGHT}[align]
    pf.first_line_indent = Cm(indent)
    pf.left_indent = Cm(left)
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.line_spacing = line
    pf.keep_with_next = keep_with_next
    add_runs(p, text, bold=bold, italic=italic, size=size)
    return p


def add_table(doc, rows):
    header, body = rows[0], rows[1:]
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = t.rows[0].cells
    for i, h in enumerate(header):
        hdr[i].text = ""
        p = hdr[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(2)
        add_runs(p, h, bold=True, size=10.5)
        shade(hdr[i])
        set_cell_border(hdr[i])
    for row in body:
        cells = t.add_row().cells
        for i, val in enumerate(row[: len(header)]):
            cells[i].text = ""
            p = cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_after = Pt(2)
            add_runs(p, val, size=10.5)
            set_cell_border(cells[i])
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def convert(src, dst):
    raw = open(src, encoding="utf-8").read()
    if "<!-- DOCX START -->" in raw:
        raw = raw.split("<!-- DOCX START -->", 1)[1]
        raw = raw.split("<!-- DOCX END -->", 1)[0]
    lines = raw.split("\n")

    doc = Document()
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
    sec.left_margin, sec.right_margin = Cm(3.0), Cm(1.5)
    sec.top_margin, sec.bottom_margin = Cm(2.0), Cm(2.0)

    style = doc.styles["Normal"]
    style.font.name = BODY_FONT
    style.font.size = Pt(12)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)

    align_ctx = None
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        stripped = line.strip()

        if stripped.startswith('<div align='):
            align_ctx = "center" if 'center' in stripped else "right"
            i += 1
            continue
        if stripped == "</div>":
            align_ctx = None
            i += 1
            continue

        if not stripped:
            i += 1
            continue

        if stripped in ("---", "***", "___"):
            i += 1
            continue

        if stripped.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c):
                    rows.append(cells)
                i += 1
            if rows:
                add_table(doc, rows)
            continue

        m = re.match(r"^(#{1,4})\s+(.*)$", stripped)
        if m:
            level, text = len(m.group(1)), m.group(2)
            size = {1: 15, 2: 13, 3: 12, 4: 12}[level]
            centered = (level == 1 or align_ctx == "center"
                        or text.startswith(("РАЗДЕЛ", "О ", "ПРИЛОЖЕНИЕ", "Приложение")))
            para(doc, text, size=size, bold=True,
                 align="center" if centered else "left",
                 indent=0 if level <= 2 else 0,
                 space_before=10 if level <= 2 else 8, space_after=6,
                 keep_with_next=True)
            i += 1
            continue

        if stripped.startswith(">"):
            text = stripped.lstrip(">").strip()
            p = para(doc, text, left=1.0, indent=0, space_after=4)
            pPr = p._p.get_or_add_pPr()
            bdr = OxmlElement("w:pBdr")
            left_el = OxmlElement("w:left")
            left_el.set(qn("w:val"), "single")
            left_el.set(qn("w:sz"), "12")
            left_el.set(qn("w:space"), "6")
            left_el.set(qn("w:color"), "999999")
            bdr.append(left_el)
            pPr.append(bdr)
            i += 1
            continue

        m = re.match(r"^[-*]\s+(.*)$", stripped)
        if m:
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.left_indent = Cm(1.25)
            add_runs(p, m.group(1))
            i += 1
            continue

        m = re.match(r"^(\d+[\.\)])\s+(.*)$", stripped)
        if m:
            para(doc, f"{m.group(1)} {m.group(2)}", indent=0, left=1.0, space_after=3)
            i += 1
            continue

        buf = [stripped]
        while i + 1 < len(lines):
            nxt = lines[i + 1].strip()
            if (not nxt or nxt.startswith(("#", "|", ">", "-", "*", "<div", "</div"))
                    or nxt in ("---", "***") or re.match(r"^\d+[\.\)]\s", nxt)):
                break
            buf.append(nxt)
            i += 1
        text = " ".join(buf)
        para(doc, text,
             align=align_ctx or "justify",
             indent=0 if align_ctx else 1.25,
             space_after=6)
        i += 1

    doc.save(dst)


if __name__ == "__main__":
    for name in sys.argv[1:]:
        out = name.rsplit(".", 1)[0] + ".docx"
        convert(name, out)
        print("OK", out)

# -*- coding: utf-8 -*-
"""Проверка DOCX: шапка, таблица, отсутствие служебных маркеров."""
import sys, os
from docx import Document
from docx.oxml.ns import qn

def cell_text(tc):
    out=[]
    for wp in tc.findall(qn("w:p")):
        out.append("".join(n.text or "" for n in wp.iter(qn("w:t"))))
    return "\n".join(out)

def check(path):
    d = Document(path)
    probs=[]
    ps=[p.text for p in d.paragraphs]
    if not any(p.startswith("ACT.") for p in ps[:3]):
        probs.append("нет шапки ACT. NNN")
    if "SECTION 1. СУТЬ ПРАВКИ." not in ps:
        probs.append("нет SECTION 1")
    for i,t in enumerate(d.tables):
        for ri,tr in enumerate(t._tbl.findall(qn("w:tr"))):
            for ci,tc in enumerate(tr.findall(qn("w:tc"))):
                txt=cell_text(tc)
                if "[[" in txt or "]]" in txt:
                    probs.append("маркер в таблице %d r%d c%d" % (i,ri,ci))
                if not txt.strip():
                    probs.append("пустая ячейка %d r%d c%d" % (i,ri,ci))
    return probs

if __name__ == "__main__":
    for p in sys.argv[1:]:
        print(os.path.basename(p), "->", check(p))

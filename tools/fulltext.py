# -*- coding: utf-8 -*-
"""Сводный документ: каждая правка — полный текст статьи (или иной единицы)."""
import copy, json, os, re
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.table import _Cell

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
BASE = os.path.join(ROOT, "Законодательная база")
TEMPLATE = os.path.join(ROOT, "Законопроекты", "ACT. 230 — Технические исправления в Закон о Национальной Гвардии.docx")
OUT = os.path.join(ROOT, "Законопроекты")
RED, GREEN = "F4CCCC", "D9EAD3"

HEADS = re.compile(r'^\s*(Статья|Глава|Раздел|Приложение|ГЛАВА)\b')
ART = re.compile(r'^\s*Статья\s*\.?\s*(\d+(?:\.\d+)*)(?!\d)(?!\.\d)')

_cache = {}
def resolve(f):
    p = os.path.join(ROOT, f)
    if os.path.exists(p):
        return p
    import glob
    key = [w for w in f.replace(".json", "").split() if len(w) > 4][:3]
    hits = [x for d in (BASE, ROOT) for x in glob.glob(os.path.join(d, "*.json")) if all(w in x for w in key)]
    if len(hits) > 1:
        pref = [x for x in hits if "Сан-Андреас" in x or "San-Andreas" in x]
        if len(pref) == 1:
            hits = pref
    if len(hits) != 1:
        raise FileNotFoundError((f, hits))
    return hits[0]

def sections(f):
    if f not in _cache:
        d = json.load(open(resolve(f), encoding="utf-8"))
        _cache[f] = [(s["n"], (s["text"] or "")) for s in d["sections"]]
    return _cache[f]

def article(f, num):
    secs = sections(f)
    for i, (n, t) in enumerate(secs):
        m = ART.match(t)
        if m and m.group(1) == num:
            buf = [t]
            for n2, t2 in secs[i+1:]:
                if HEADS.match(t2) or t2.strip() == "\u200b":
                    break
                buf.append(t2)
            return "\n".join(x.rstrip() for x in buf)
    raise KeyError((f, num))

def chapter_line(f, roman):
    for n, t in sections(f):
        if re.match(r'^\s*Глава\s+' + roman + r'\b', t):
            return t.split("\n")[0].strip()
    raise KeyError((f, roman))

def mark(text, pairs):
    left, right = text, text
    for old, new in pairs:
        c = left.count(old)
        assert c == 1, "не уникально (%d): %r" % (c, old[:70])
        left = left.replace(old, "[[R]]%s[[/R]]" % old)
    for old, new in pairs:
        c = right.count(old)
        assert c == 1, "не уникально (%d): %r" % (c, old[:70])
        right = right.replace(old, "[[G]]%s[[/G]]" % new)
    return left, right

def row(f, num, pairs):
    return mark(article(f, num), pairs)

# ---------- отрисовка ----------
def fill_tc(tc, text, tr):
    for p in tc.findall(qn("w:p")):
        tc.remove(p)
    paras = [[("", None)]]
    tokens, pos = [], 0
    for m in re.finditer(r"\[\[(R|G)\]\](.*?)\[\[/\1\]\]", text, re.S):
        if m.start() > pos:
            tokens.append((text[pos:m.start()], None))
        tokens.append((m.group(2), RED if m.group(1) == "R" else GREEN))
        pos = m.end()
    if pos < len(text):
        tokens.append((text[pos:], None))
    for seg, fill in tokens:
        parts = seg.split("\n")
        for i, part in enumerate(parts):
            if part:
                paras[-1].append((part, fill))
            if i < len(parts) - 1:
                paras.append([("", None)])
    for runs in paras:
        p = OxmlElement("w:p")
        pPr = OxmlElement("w:pPr")
        wc = OxmlElement("w:widowControl"); wc.set(qn("w:val"), "0"); pPr.append(wc)
        sp = OxmlElement("w:spacing"); sp.set(qn("w:line"), "240"); sp.set(qn("w:lineRule"), "auto"); pPr.append(sp)
        jc = OxmlElement("w:jc"); jc.set(qn("w:val"), "both"); pPr.append(jc)
        p.append(pPr)
        for seg, fill in runs:
            if not seg:
                continue
            r = OxmlElement("w:r")
            rPr = OxmlElement("w:rPr")
            rf = OxmlElement("w:rFonts")
            for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
                rf.set(qn(a), "Times New Roman")
            rPr.append(rf)
            col = OxmlElement("w:color"); col.set(qn("w:val"), "333333"); rPr.append(col)
            sz = OxmlElement("w:sz"); sz.set(qn("w:val"), "24"); rPr.append(sz)
            szcs = OxmlElement("w:szCs"); szcs.set(qn("w:val"), "24"); rPr.append(szcs)
            if fill:
                shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear")
                shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), fill); rPr.append(shd)
            r.append(rPr)
            t = OxmlElement("w:t"); t.set(qn("xml:space"), "preserve"); t.text = seg
            r.append(t)
            p.append(r)
        tc.append(p)

def make_table(tbl_tpl, tr_tpl, rows):
    tbl = copy.deepcopy(tbl_tpl)
    trs = tbl.findall(qn("w:tr"))
    for tr in trs[1:]:
        tbl.remove(tr)
    for left, right in rows:
        tr = copy.deepcopy(tr_tpl)
        tbl.append(tr)
        tcs = tr.findall(qn("w:tc"))
        fill_tc(tcs[0], left, tr)
        fill_tc(tcs[1], right, tr)
    return tbl

def build(groups, out_name, head, subtitle, section1):
    doc = Document(TEMPLATE)
    body = doc.element.body
    ps = doc.paragraphs
    ps[0].runs[0].text = head
    for r in ps[0].runs[1:]:
        r._r.getparent().remove(r._r)
    ps[3].runs[0].text = "OCTOBER 21, 2026"
    for r in ps[3].runs[1:]:
        r._r.getparent().remove(r._r)
    ps[6].runs[0].text = subtitle
    for r in ps[6].runs[1:]:
        r._r.getparent().remove(r._r)
    ps[9].runs[0].text = section1
    for r in ps[9].runs[1:]:
        r._r.getparent().remove(r._r)

    p_sec2 = ps[10]._p
    tbl_tpl = doc.tables[0]._tbl
    tr_tpl = copy.deepcopy(doc.tables[0].rows[1]._tr)
    p_after = ps[11]._p

    for el in (p_sec2, tbl_tpl):
        body.remove(el)

    anchor = p_after
    for title, rows in groups:
        ph = copy.deepcopy(p_sec2)
        for r in ph.findall(qn("w:r")):
            ph.remove(r)
        src = p_sec2.findall(qn("w:r"))
        if src:
            r = copy.deepcopy(src[0])
            for t in r.findall(qn("w:t")):
                t.text = title
            ph.append(r)
        anchor.addprevious(ph)
        t = make_table(tbl_tpl, tr_tpl, rows)
        anchor.addprevious(t)

    path = os.path.join(OUT, out_name)
    doc.save(path)
    print("OK", out_name, "| групп:", len(groups), "| строк:", sum(len(r) for _, r in groups))
    return path

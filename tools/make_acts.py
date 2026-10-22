# -*- coding: utf-8 -*-
"""Генератор законопроектов по образцу ACT. 230.
Клонирует оформление шаблона и заменяет содержательную часть.
"""
import copy, re, sys, os
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.table import _Cell

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
TEMPLATE = os.path.join(ROOT, "ACT. 230 — Технические исправления в Закон о Национальной Гвардии.docx")
RED, GREEN = "F4CCCC", "D9EAD3"
TOKEN = re.compile(r"\[\[(R|G)\]\](.*?)\[\[/\1\]\]", re.S)


def parse_line(text):
    segs, pos = [], 0
    for m in TOKEN.finditer(text):
        if m.start() > pos:
            segs.append((text[pos:m.start()], None))
        segs.append((m.group(2), RED if m.group(1) == "R" else GREEN))
        pos = m.end()
    if pos < len(text):
        segs.append((text[pos:], None))
    return segs


def add_run(p, text, fill=None):
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    rFonts = r._r.rPr.find(qn("w:rFonts"))
    for attr in ("w:eastAsia", "w:hAnsi", "w:cs"):
        rFonts.set(qn(attr), "Times New Roman")
    r.font.size = Pt(12)
    r.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
    rPr = r._r.get_or_add_rPr()
    szCs = OxmlElement("w:szCs")
    szCs.set(qn("w:val"), "24")
    rPr.append(szCs)
    if fill:
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), fill)
        r._r.get_or_add_rPr().append(shd)


def fill_cell(cell, text):
    for p in list(cell.paragraphs):
        p._p.getparent().remove(p._p)
    for line in text.split("\n"):
        p = cell.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.widow_control = False
        p.paragraph_format.line_spacing = 1.0
        for seg, fill in parse_line(line):
            if seg:
                add_run(p, seg, fill)


def set_par_text(p, text):
    runs = p.runs
    runs[0].text = text
    for r in runs[1:]:
        r._r.getparent().remove(r._r)


def build(bill, outdir="."):
    doc = Document(TEMPLATE)
    ps = doc.paragraphs
    set_par_text(ps[0], f"ACT. {bill['num']}")
    set_par_text(ps[3], bill["date"])
    set_par_text(ps[6], bill["subtitle"])
    set_par_text(ps[9], bill["section1"])
    set_par_text(ps[10], bill["section2"])

    tbl = doc.tables[0]
    template_tr = copy.deepcopy(tbl.rows[1]._tr)
    for row in list(tbl.rows)[1:]:
        tbl._tbl.remove(row._tr)

    for left, right in bill["rows"]:
        tr = copy.deepcopy(template_tr)
        tbl._tbl.append(tr)
        cells = tr.findall(qn("w:tc"))
        fill_cell(_Cell(cells[0], tbl), left)
        fill_cell(_Cell(cells[1], tbl), right)

    doc.save(f"{outdir}/{bill['file']}")
    print("OK", bill["file"])
